#!/usr/bin/env python3
"""Prépare 4 lots de relance de 400 contacts (mercredi → samedi) pour le salon du 5 octobre.

Source unique : feuille « Tous les contacts » du fichier Master.

Exclusions : désabonnés, inscrits (statut, colonne Inscrit, fichier d'inscriptions), e-mails envoyés
lundi 28/09 et mardi 29/09, catégories hors cible, contacts jamais envoyés (Nb envois = 0).

Priorités (validées avec l'utilisateur le 30/09) :
  1 Cavistes, bars à vin, caves à manger, fromageries-caves, torréfacteurs
  2 Épiceries fines, coffrets gourmands, magasins de bouche fins
  3 Tables bistronomiques, restaurants gastronomiques, hôtels-restaurants
  4 Traiteurs réception, lieux de réception
  5 Autres restaurants et CHR (les moins sollicités d'abord)
Contacts sans catégorie : profil déduit du nom de l'établissement (formule « à toute l'équipe de… »)
et de la campagne (« Vague Cavistes »). PMU, pizzerias, crêperies, kebabs et restauration rapide exclus.

Usage : python generer_relances.py [chemin/vers/Master.xlsx]
"""
import csv
import glob
import os
import re
import sys
import unicodedata
from collections import Counter
from difflib import SequenceMatcher

import openpyxl

ICI = os.path.dirname(os.path.abspath(__file__))
SOURCES = os.path.join(ICI, "donnees_sources")
TAILLE_LOT = 400
LOTS = ["RELANCE_LOT1_MERCREDI.csv", "RELANCE_LOT2_JEUDI.csv",
        "RELANCE_LOT3_VENDREDI.csv", "RELANCE_LOT4_SAMEDI.csv"]

CATEGORIES_EXCLUES = {
    # hors cible, 0 % de conversion
    "Producteur - Vente Directe", "Camping", "Camping Premium", "Gîte", "Gîte de Groupe",
    "Chambre d'hôtes", "Maison d'Hôtes",
    # producteurs et distributeurs
    "Brasserie Artisanale", "Œnotourisme", "Magasin de Producteurs", "Distributeur CHR",
    # entreprises, réseaux et événementiel (hors hiérarchie)
    "Réseau de Dirigeants", "Club d'Affaires", "Agence Événementielle", "Prestataire Événementiel",
    "Cadeaux d'Affaires", "Golf / Resort", "Conciergerie d'Entreprise", "CSE / Comité d'Entreprise",
    "Wedding Planner", "Salle de Séminaire", "Club Sportif / Partenaires", "Espace Coworking",
    "Organisateur de Congrès", "Autocariste", "Site de Loisirs", "École Hôtelière",
    "Agence de Développement", "Thalasso-Spa", "Chef à Domicile",
}
PRIO_CATEGORIE = {
    "Caviste": 1, "Bar à Vin": 1, "Bar à Vin / Bar-restaurant": 1, "Fromagerie-Cave": 1, "Torréfacteur": 1,
    "Épicerie Fine": 2, "Coffrets Gourmands": 2, "Chocolaterie-Coffrets": 2, "Poissonnerie-Écailler": 2,
    "Table Bistronomique": 3, "Restaurant Gastronomique": 3, "Hôtel-Restaurant": 3, "Hôtels-restaurants": 3,
    "Traiteur Réception": 4, "Lieu de Réception": 4, "Traiteur": 4, "Boucherie-Traiteur": 4,
    "Restaurant": 5, "Ferme-Auberge": 5, "Casino-Restaurant": 5,
}
# Profil déduit du nom (contacts sans catégorie), dans l'ordre de test.
REGLES_NOM = [
    (1, "Caviste / bar à vin (déduit)",
     r"\bcaves?\b|caviste|\bvins?\b|vinoth|\bchais?\b|cellier|tapas|fromag|torr[ée]f|\bbib[oe]"),
    (2, "Épicerie fine (déduit)", r"[ée]picer|comptoir gourmand|coffret|chocolat"),
    (4, "Traiteur (déduit)", r"traiteur|r[ée]ception"),
    (3, "Bistrot / gastronomique / hôtel-restaurant (déduit)",
     r"bistr|gastronom|\bh[oô]tel|auberge|relais|logis|\btable\b"),
]
RE_RAPIDE = re.compile(r"\bpmu\b|tabac|pizz|cr[eê]p|galette|kebab|burger|tacos|sushi|snack|friterie|"
                       r"fast.?food|restauration rapide|sandwich|bagel|wok|roi du poulet|diner\b|tommy", re.I)
# Règles de rejet d'origine (brief initial) : écoles, campings, gîtes, chambres d'hôtes…
RE_HORS_CIBLE = re.compile(r"lyc[ée]e|[ée]cole|coll[èe]ge|\bcfa\b|greta|campus|camping|g[iî]tes?\b|"
                           r"chambres? d.h[oô]tes?|maison d.h[oô]tes|ehpad|scolaire", re.I)
SOUS_CHAINES_INTERDITES = ("le21pornic", "21pornic", "domainedeliziec", "liziec", "camping", "gite", "greta", "cfa", "chambre-hote",
                           "ac-nantes.fr", "ac-rennes.fr")
RE_EMAIL = re.compile(r"[a-z0-9._%+'-]+@[a-z0-9.-]+\.[a-z]{2,}")
WEBMAILS = {"gmail.com", "orange.fr", "wanadoo.fr", "free.fr", "sfr.fr", "laposte.net", "hotmail.com",
            "hotmail.fr", "outlook.fr", "outlook.com", "yahoo.fr", "yahoo.com", "icloud.com", "bbox.fr",
            "live.fr", "aol.com", "neuf.fr", "club-internet.fr", "numericable.fr", "gmx.fr", "msn.com"}
# Inscrits signalés par l'utilisateur le 30/09 (adresses alternatives d'établissements déjà inscrits).
INSCRITS_SIGNALES = {"bulleblancrouge@gmail.com", "ralph.pillet@wanadoo.fr", "ecrivez-nous@larayonnantes.fr",
                     "ginoalteregos@gmail.com", "dv-cave@dvfrance.com", "c.ribet@dvfrance.com",
                     "comptabilitefournisseur@dvfrance.com", "n.vergnault@dvfrance.com"}
MOTS_VIDES = {"le", "la", "les", "l", "de", "du", "des", "d", "et", "sarl", "sas", "scea", "eurl", "sa",
              "societe", "ste", "entreprise", "restaurant", "a", "au", "aux", "the", "sarl", "dv"}


def norm(e):
    return (e or "").strip().lower().removeprefix("mailto:").strip(" ;,.")


def trouver(motif):
    for dossier in (SOURCES, ICI):
        f = glob.glob(os.path.join(dossier, motif))
        if f:
            return max(f, key=os.path.getmtime)  # la version la plus récente
    return None


def emails_csv(chemin):
    with open(chemin, encoding="utf-8-sig") as f:
        return {norm(r.get("Email")) for r in csv.DictReader(f)} - {""}


def cle_nom(nom):
    """Nom d'établissement réduit à ses mots utiles, sans accents ni ponctuation."""
    n = unicodedata.normalize("NFD", (nom or "").lower())
    n = "".join(c for c in n if not unicodedata.combining(c))
    n = re.sub(r"^a toute l'equipe\b", " ", n.replace("’", "'"))
    mots = [m for m in re.split(r"[^a-z0-9]+", n) if m and m not in MOTS_VIDES]
    return " ".join(mots)


def domaine_pro(e):
    d = e.rsplit("@", 1)[-1] if "@" in e else ""
    return d if d and d not in WEBMAILS else ""


def lire_inscriptions(chemin):
    """E-mails, domaines pros, parties avant @ et noms d'établissements des inscrits."""
    wb = openpyxl.load_workbook(chemin, read_only=True)
    ins = {"emails": set(), "domaines": set(), "locaux": set(), "noms": set()}
    for ws in wb.worksheets:
        lignes = list(ws.iter_rows(values_only=True))
        if not lignes:
            continue
        titres = [str(t or "").lower() for t in lignes[0]]
        for l in lignes[1:]:
            for j, t in enumerate(titres):
                v = str(l[j] or "").strip() if j < len(l) else ""
                if not v:
                    continue
                if "mail" in t:
                    e = norm(v.replace(" gmail.com", "@gmail.com"))
                    ins["emails"].add(e)
                    ins["locaux"].add(e.split("@")[0])
                    if domaine_pro(e):
                        ins["domaines"].add(domaine_pro(e))
                elif "site" in t:
                    h = re.sub(r"^(https?://)?(www\.)?", "", v.lower()).split("/")[0]
                    if "." in h and not re.search(r"facebook|instagram|google|linkedin", h):
                        ins["domaines"].add(h)
                elif t.startswith("nom de la soci"):
                    if cle_nom(v):
                        ins["noms"].add(cle_nom(v))
    ins["locaux"] -= {"contact", "info", "bonjour", "hello", "accueil", "reservation", "direction",
                       "restaurant", "cave", "commande", "vin", "boutique", "infos"}
    return ins


def meme_inscrit(email, prenom, etablissement, ins):
    """Raison si le contact correspond à un établissement inscrit, sinon None."""
    if email in ins["emails"] or email in INSCRITS_SIGNALES:
        return "e-mail inscrit"
    if domaine_pro(email) in ins["domaines"]:
        return "même domaine qu'un inscrit"
    if email.split("@")[0] in ins["locaux"]:
        return "même adresse qu'un inscrit, autre fournisseur"
    for nom in (etablissement, prenom if prenom.lower().startswith("à toute l'équipe") else ""):
        cle = cle_nom(nom)
        if not cle:
            continue
        mots = set(cle.split())
        for n in ins["noms"]:
            mots_n = set(n.split())
            if cle == n or SequenceMatcher(None, cle, n).ratio() >= 0.9 \
                    or (len(mots_n) >= 2 and mots_n <= mots) or (len(mots) >= 2 and mots <= mots_n) \
                    or (len(mots_n) >= 2 and len(mots_n & mots) >= 2):
                return f"même établissement qu'un inscrit ({n})"
    return None


def profil(categorie, prenom, campagnes, email):
    """(priorité, libellé) ou (None, raison d'exclusion)."""
    if categorie:
        if RE_HORS_CIBLE.search(prenom if prenom.lower().startswith("à toute l'équipe") else "") or \
                any(x in email for x in SOUS_CHAINES_INTERDITES):
            return None, "école, camping, gîte ou chambre d'hôtes"
        if categorie in CATEGORIES_EXCLUES:
            return None, f"catégorie exclue : {categorie}"
        return PRIO_CATEGORIE.get(categorie, 5), categorie
    nom = prenom if prenom.lower().startswith("à toute l'équipe") else ""
    if RE_HORS_CIBLE.search(nom) or any(x in email for x in SOUS_CHAINES_INTERDITES):
        return None, "école, camping, gîte ou chambre d'hôtes"
    if RE_RAPIDE.search(nom) or RE_RAPIDE.search(email.split("@")[0]):
        return None, "PMU / pizzeria / restauration rapide"
    if "Vague Cavistes" in campagnes:
        return 1, "Caviste (campagne Cavistes)"
    for prio, libelle, motif in REGLES_NOM:
        if re.search(motif, nom + " " + email.split("@")[0], re.I):
            return prio, libelle
    return 5, "Restaurant / CHR (sans catégorie)"


def main():
    master = sys.argv[1] if len(sys.argv) > 1 else trouver("*Master_tous_les_contacts*.xlsx")
    lundi = trouver("*CSV_ENVOI_LUNDI_4_LOTS_1*.csv")
    mardi = os.path.join(ICI, "NOUVEAUX_CONTACTS_CLAUDE_CODE_TOUS.csv")
    inscriptions = trouver("*Inscription*.xlsx")
    for nom, chemin in (("Master", master), ("Envoi lundi", lundi), ("Envoi mardi", mardi),
                        ("Inscriptions", inscriptions)):
        if not chemin or not os.path.exists(chemin):
            sys.exit(f"Fichier introuvable : {nom}")
    deja_envoyes = emails_csv(lundi) | emails_csv(mardi)
    inscrits = lire_inscriptions(inscriptions)
    doublons_inscrits = []

    ws = openpyxl.load_workbook(master, read_only=True)["Tous les contacts"]
    lignes = list(ws.iter_rows(values_only=True))
    ix = {k: i for i, k in enumerate(lignes[0])}
    rejets, candidats, vus = Counter(), [], set()
    for l in lignes[1:]:
        email = norm(l[ix["Email"]])
        statut = l[ix["Statut"]] or ""
        n = int(l[ix["Nb envois"]] or 0)
        prenom = (l[ix["Prénom"]] or "").strip() or "à toute l'équipe"
        if not RE_EMAIL.fullmatch(email):
            rejets["e-mail invalide"] += 1
        elif email in vus:
            rejets["doublon"] += 1
        elif "DÉSABONNÉ" in statut.upper():
            rejets["désabonné"] += 1
        elif "INSCRIT" in statut.upper() or l[ix["Inscrit"]]:
            rejets["inscrit au salon"] += 1
        elif meme_inscrit(email, prenom, l[ix["Établissement"]], inscrits):
            rejets["inscrit au salon (autre adresse ou même établissement)"] += 1
            doublons_inscrits.append((email, prenom, meme_inscrit(email, prenom, l[ix["Établissement"]], inscrits)))
        elif email in deja_envoyes:
            rejets["déjà envoyé lundi ou mardi"] += 1
        elif n < 1:
            rejets["jamais envoyé (Nb envois = 0)"] += 1
        else:
            prio, libelle = profil(l[ix["Catégorie"]], prenom, l[ix["Campagnes"]] or "", email)
            if prio is None:
                rejets[libelle if not libelle.startswith("catégorie") else "catégorie hors cible"] += 1
            else:
                candidats.append({"email": email, "prenom": prenom, "prio": prio, "profil": libelle,
                                  "envois": n})
        vus.add(email)

    # Meilleure priorité d'abord, puis les moins sollicités ; ordre du Master à égalité.
    candidats.sort(key=lambda c: (c["prio"], c["envois"]))
    total = TAILLE_LOT * len(LOTS)
    if len(candidats) < total:
        sys.exit(f"Seulement {len(candidats)} contacts éligibles pour {total} demandés.")
    retenus = candidats[:total]

    print(f"Master : {len(lignes) - 1} contacts · éligibles : {len(candidats)} · retenus : {total}")
    for raison, nb in rejets.most_common():
        print(f"  écartés — {raison} : {nb}")
    for i, fichier in enumerate(LOTS):
        lot = retenus[i * TAILLE_LOT:(i + 1) * TAILLE_LOT]
        with open(os.path.join(ICI, fichier), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(["Email", "Prenom"])
            w.writerows((c["email"], c["prenom"]) for c in lot)
        print(f"\n## {fichier} — {len(lot)} contacts")
        print("  Priorités :", ", ".join(f"P{p} = {n}" for p, n in sorted(Counter(c['prio'] for c in lot).items())))
        print("  Envois précédents :", ", ".join(f"{e} envoi{'s' if e > 1 else ''} = {n}"
                                               for e, n in sorted(Counter(c['envois'] for c in lot).items())))
        for p, n in Counter(c["profil"] for c in lot).most_common():
            print(f"    {n:4d}  {p}")
    print("\nInscrits retrouvés sous une autre adresse ou un nom identique :")
    for e, p, r in doublons_inscrits:
        print(f"  {e:45s} {p[:45]:45s} {r}")
    with open(os.path.join(ICI, "RELANCE_DETAILS.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["Lot", "Email", "Prenom", "Priorite", "Profil", "Nb_envois_precedents"])
        for i, c in enumerate(retenus):
            w.writerow([LOTS[i // TAILLE_LOT], c["email"], c["prenom"], c["prio"], c["profil"], c["envois"]])


if __name__ == "__main__":
    main()
